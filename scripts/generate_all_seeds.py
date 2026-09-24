"""한우리_초기데이터_이행.xlsx → 3개 시드(kt) 재생성.

시트 구성:
  - 장부내역: 통장 실입출금(잔액 SSOT). 회비/경조 도메인과 별개.
  - 회비내역: 회원별 반기 회비(도메인 테이블). 장부에 중복 삽입 안 함.
  - 경조사비: 경조 이력(도메인 테이블). 장부에 중복 삽입 안 함.

규칙:
  - 잔액 = 장부내역 입금합 - 출금합 (인위적 보정 없음).
"""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

EXCEL_PATH = Path(r"C:\Users\gyoum\Downloads\한우리_초기데이터_이행.xlsx")
SEED_DIR = (
    Path(__file__).resolve().parents[1]
    / "app/src/main/kotlin/com/smartexpense/data/local/seed"
)

# --- 카테고리 (ClubEnums.ClubCategory 기준) ---
REGULAR_DUES = "정기 회비 (월/연회비)"
DONATION = "찬조금/특별 회비 (기부금 등)"
INTEREST_INCOME = "이자/수익"
OTHER_INCOME = "기타 수입 (이월금 등)"
MEAL = "식대/다과비 (모임 식사, 카페 등)"
VENUE_RENTAL = "장소 대관료 (모임 장소 대여비)"
EVENT = "행사/진행비 (이벤트 기획, 상품 구매 등)"
SUPPLIES = "비품/소모품비 (운영에 필요한 물품 구매)"
CONDOLENCE = "경조사비 (회원 화환, 축의금/조의금 등)"
TRANSPORT = "교통/통신비 (이동 경비, 문자 발송비 등)"
OTHER_EXPENSE = "기타 지출 (수수료 및 예비비)"


def norm_date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if re.match(r"\d{4}-\d{2}-\d{2}", text):
        return text[:10]
    if re.match(r"\d{4}\.\d{2}\.\d{2}", text):
        return text.replace(".", "-")
    return ""


def norm_amount(value) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).replace(",", "").strip()
    if not text or text in ("-", "—"):
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def kotlin_escape(text: str) -> str:
    return (
        text.replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("\r\n", "\\n")
        .replace("\n", "\\n")
        .replace("\r", "\\n")
    )


# =========================================================================
# 1) 장부내역 → TransactionInitialData.kt
# =========================================================================
BANK_NAMES = ("하나은행", "우체국", "농협", "신한", "국민은행", "은행")


def classify(label: str, income: int, expense: int, note: str) -> tuple[str, str]:
    text = f"{label} {note}".lower()
    if income > 0:
        if any(k in text for k in ("이자",)):
            return "INCOME", INTEREST_INCOME
        if any(k in text for k in ("찬조", "기부", "후원")):
            return "INCOME", DONATION
        if any(k in text for k in ("회비", "분납", "납부", "미납", "입회비")):
            return "INCOME", REGULAR_DUES
        if any(k in text for k in ("이월", "잔여", "잔액", "전대", "반환", "환급", "되돌려")):
            return "INCOME", OTHER_INCOME
        if any(k in (label) for k in BANK_NAMES):
            return "INCOME", INTEREST_INCOME
        if re.fullmatch(r"[가-힣]{2,5}", label.strip()):
            return "INCOME", REGULAR_DUES
        return "INCOME", OTHER_INCOME

    # 지출
    if any(k in text for k in ("경조", "화환", "꽃", "칠순", "고희", "회갑", "조의", "축의",
                               "부의", "출산", "결혼", "돌잔치", "개업화환", "근조", "상조", "부고")):
        return "EXPENSE", CONDOLENCE
    if any(k in text for k in ("펜션", "숙박", "숙소", "대관", "모텔", "리조트")):
        return "EXPENSE", VENUE_RENTAL
    if any(k in text for k in ("모임", "식대", "식비", "다과", "부식", "주류", "음료", "가리비",
                               "케잌", "케이크", "정산", "바베큐", "회식", "식사", "명절", "저녁", "아침")):
        return "EXPENSE", MEAL
    if any(k in text for k in ("행사", "이벤트", "진행비", "상품", "게임", "경품", "레크레이션")):
        return "EXPENSE", EVENT
    if any(k in text for k in ("비품", "소모품", "운영비", "잡비", "용품", "구입비")):
        return "EXPENSE", SUPPLIES
    if any(k in text for k in ("교통", "통신", "문자", "주유", "유류", "택시", "톨비",
                               "송금수수료", "수수료", "통행", "주차")):
        return "EXPENSE", TRANSPORT
    return "EXPENSE", OTHER_EXPENSE


def build_note(label: str, remark: str) -> str:
    remark = (remark or "").strip()
    label = (label or "").strip()
    combined = f"{label} | {remark}" if remark and remark != label else label
    return " ".join(combined.split())


def gen_transactions(wb) -> int:
    ws = wb["장부내역"]
    records = []
    running = 0
    for row in ws.iter_rows(min_row=2, values_only=True):
        label = str(row[0]).strip() if row[0] else ""
        pay_date = norm_date(row[1])
        income = norm_amount(row[2])
        expense = norm_amount(row[3])
        remark = str(row[4]).strip() if len(row) > 4 and row[4] else ""
        if income == 0 and expense == 0:
            continue
        if not pay_date:
            continue
        tx_type, category = classify(label, income, expense, remark)
        running += income - expense
        records.append({
            "date": pay_date,
            "type": tx_type,
            "category": category,
            "income": income,
            "expense": expense,
            "note": build_note(label, remark),
            "balance": running,
        })

    lines = [
        "package com.smartexpense.data.local.seed",
        "",
        "import com.smartexpense.data.local.entity.club.ClubTransactionType",
        "",
        "/**",
        " * 한우리 입출금 기록부(통장) 시드 — 엑셀 `한우리_초기데이터_이행.xlsx` [장부내역] 시트.",
        " * 잔액 SSOT: 최종 잔액 = 전체 수입 합 − 전체 지출 합.",
        " * 회비/경조 도메인 이력은 각각 DuesInitialData / EventExpenseInitialData 로 분리 관리하며",
        " * 이 장부에 중복 삽입하지 않는다.",
        " */",
        "object TransactionInitialData {",
        "",
        "    data class TransactionSeedRecord(",
        "        val date: String,",
        "        val type: ClubTransactionType,",
        "        val category: String,",
        "        val incomeAmount: Int,",
        "        val expenseAmount: Int,",
        "        val note: String,",
        "        val balanceAfter: Int? = null,",
        "    )",
        "",
        "    val records: List<TransactionSeedRecord> = listOf(",
    ]
    body = []
    for r in records:
        note = kotlin_escape(r["note"])
        body.append(
            f'        TransactionSeedRecord("{r["date"]}", '
            f'ClubTransactionType.{r["type"]}, "{r["category"]}", '
            f'{r["income"]}, {r["expense"]}, "{note}", {r["balance"]})'
        )
    lines.append(",\n".join(body))
    lines += ["    )", "}", ""]
    (SEED_DIR / "TransactionInitialData.kt").write_text("\n".join(lines), encoding="utf-8")

    inc = sum(r["income"] for r in records)
    exp = sum(r["expense"] for r in records)
    print(f"[장부내역] {len(records)}건  수입 {inc:,}  지출 {exp:,}  잔액 {inc - exp:,}")
    return inc - exp


# =========================================================================
# 2) 회비내역 → DuesInitialData.kt
# =========================================================================
STANDARD_TERM_AMOUNT = 150000


def split_cell(value) -> list:
    """셀 하나에 줄바꿈으로 여러 값이 들어있는 경우(분납)를 분해."""
    if value is None:
        return []
    if isinstance(value, (datetime, date, int, float)):
        return [value]
    text = str(value).strip()
    if not text:
        return []
    return [p.strip() for p in re.split(r"[\n\r]+", text) if p.strip()]


def gen_dues(wb) -> None:
    ws = wb["회비내역"]
    members: list[tuple[str, list[dict]]] = []
    state = {"name": None, "records": [], "year": None, "group": None}

    def flush_group() -> None:
        g = state["group"]
        if g is None:
            return
        dates = g["dates"]
        amounts = g["amounts"]
        # 금액셀 병합(날짜 수 > 금액 수): 총액을 날짜 수만큼 균등 분할
        if len(amounts) < len(dates) and amounts:
            total = sum(amounts)
            n = len(dates)
            base = total // n
            payments = []
            for i, d in enumerate(dates):
                a = base + (total - base * n if i == 0 else 0)
                payments.append((d, a))
        else:
            payments = []
            for i, a in enumerate(amounts):
                d = dates[i] if i < len(dates) else (dates[-1] if dates else "")
                if a > 0 and d:
                    payments.append((d, a))

        if not payments:
            rec = {"year": g["year"], "term": g["term"],
                   "amount": STANDARD_TERM_AMOUNT, "payDate": "",
                   "isPaid": False, "payments": None}
        elif len(payments) == 1:
            d, a = payments[0]
            rec = {"year": g["year"], "term": g["term"],
                   "amount": a, "payDate": d, "isPaid": True, "payments": None}
        else:
            paid_sum = sum(a for _, a in payments)
            rec = {"year": g["year"], "term": g["term"],
                   "amount": STANDARD_TERM_AMOUNT, "payDate": "",
                   "isPaid": paid_sum >= STANDARD_TERM_AMOUNT, "payments": payments}
        state["records"].append(rec)
        state["group"] = None

    def flush_member() -> None:
        flush_group()
        if state["name"] is not None:
            members.append((state["name"], state["records"]))

    for row in ws.iter_rows(min_row=2, values_only=True):
        mem, year_s, term, pdate, amt = row[0], row[1], row[2], row[3], row[4]
        if mem:
            flush_member()
            state["name"] = str(mem).strip()
            state["records"] = []
            state["year"] = None
        if year_s:
            state["year"] = int(re.sub(r"[^0-9]", "", str(year_s)))
        term_label = str(term).strip() if term else ""
        if term_label:  # 새 반기 그룹 시작
            flush_group()
            state["group"] = {"year": state["year"], "term": term_label,
                              "dates": [], "amounts": []}
        if state["group"] is None:
            continue
        # 이 행의 납부(분납 셀 \n 분해) 누적
        amt_cells = split_cell(amt)
        if any("미납" in str(a) for a in amt_cells):
            continue  # 미납 → 납부 없음
        for d in split_cell(pdate):
            nd = norm_date(d)
            if nd:
                state["group"]["dates"].append(nd)
        for a in amt_cells:
            na = norm_amount(a)
            if na > 0:
                state["group"]["amounts"].append(na)
    flush_member()

    lines = [
        "package com.smartexpense.data.local.seed",
        "",
        "import com.smartexpense.data.local.entity.club.DuesPaymentMethod",
        "",
        "/**",
        " * 회비 납부 이력 시드 — 엑셀 `한우리_초기데이터_이행.xlsx` [회비내역] 시트.",
        " * 도메인 테이블(yearly_dues / dues_detail / dues_payment_history) 전용.",
        " * 장부(club_transactions)에 중복 삽입하지 않는다.",
        " */",
        "object DuesInitialData {",
        "",
        "    data class DuesPaymentSeed(",
        "        val payDate: String,",
        "        val amount: Int,",
        "    )",
        "",
        "    data class DuesSeedRecord(",
        "        val year: Int,",
        "        val termLabel: String,",
        "        val amount: Int,",
        "        val payDate: String,",
        "        val isPaid: Boolean,",
        "        val paidAmount: Int? = null,",
        "        val payments: List<DuesPaymentSeed>? = null,",
        "    )",
        "",
        "    data class MemberDuesSeed(",
        "        val paymentMethod: DuesPaymentMethod,",
        "        val records: List<DuesSeedRecord>",
        "    )",
        "",
        "    val seedsByMember: Map<String, MemberDuesSeed> = mapOf(",
    ]
    mem_blocks = []
    total_records = 0
    for name, recs in members:
        rec_lines = []
        for r in recs:
            total_records += 1
            payments = r.get("payments")
            if payments:
                pay_str = ", ".join(
                    f'DuesPaymentSeed("{d}", {a})' for d, a in payments
                )
                rec_lines.append(
                    f'            DuesSeedRecord({r["year"]}, "{r["term"]}", '
                    f'{r["amount"]}, "{r["payDate"]}", {str(r["isPaid"]).lower()}, '
                    f'payments = listOf({pay_str}))'
                )
            else:
                rec_lines.append(
                    f'            DuesSeedRecord({r["year"]}, "{r["term"]}", '
                    f'{r["amount"]}, "{r["payDate"]}", {str(r["isPaid"]).lower()})'
                )
        block = (
            f'        "{name}" to MemberDuesSeed(\n'
            f'            paymentMethod = DuesPaymentMethod.HALF_YEARLY,\n'
            f'            records = listOf(\n'
            + ",\n".join(rec_lines)
            + "\n            )\n"
            f'        )'
        )
        mem_blocks.append(block)
    lines.append(",\n".join(mem_blocks))
    lines += [
        "    )",
        "",
        "    fun totalTargetAmountForYear(records: List<DuesSeedRecord>, year: Int): Int =",
        "        records.filter { it.year == year }.sumOf { it.amount }",
        "}",
        "",
    ]
    (SEED_DIR / "DuesInitialData.kt").write_text("\n".join(lines), encoding="utf-8")
    print(f"[회비내역] {len(members)}명  {total_records}건")


# =========================================================================
# 3) 경조사비 → EventExpenseInitialData.kt
# =========================================================================
def gen_events(wb) -> None:
    ws = wb["경조사비"]
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        mem, cat, sub, amt, note = row[0], row[1], row[2], row[3], row[4]
        if not mem:
            continue
        cat = str(cat).strip() if cat else ""
        sub = str(sub).strip() if sub else ""
        amount = norm_amount(amt)
        if cat == "상조":
            raw = f"상조({sub})"
        elif cat == "결혼":
            raw = "결혼"
        else:
            raw = cat
        records.append({"member": str(mem).strip(), "raw": raw, "amount": amount})

    lines = [
        "package com.smartexpense.data.local.seed",
        "",
        "/**",
        " * 경조사비 이력 시드 — 엑셀 `한우리_초기데이터_이행.xlsx` [경조사비] 시트.",
        " * 도메인 테이블 `event_expenses` 전용. 장부(club_transactions)에 넣지 않으며 통장 잔액과 무관.",
        " */",
        "object EventExpenseInitialData {",
        "",
        '    const val SEED_DATE = "2025-01-01"',
        "",
        "    data class EventExpenseSeedRecord(",
        "        val memberName: String,",
        "        val rawSubCategory: String,",
        "        val amount: Int",
        "    )",
        "",
        "    val records: List<EventExpenseSeedRecord> = listOf(",
    ]
    body = []
    for r in records:
        body.append(
            f'        EventExpenseSeedRecord("{r["member"]}", "{r["raw"]}", {r["amount"]})'
        )
    lines.append(",\n".join(body))
    lines += ["    )", "}", ""]
    (SEED_DIR / "EventExpenseInitialData.kt").write_text("\n".join(lines), encoding="utf-8")
    total = sum(r["amount"] for r in records)
    print(f"[경조사비] {len(records)}건  합계 {total:,}")


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    balance = gen_transactions(wb)
    gen_dues(wb)
    gen_events(wb)
    wb.close()
    print(f"\n최종 통장 잔액(장부 기준) = {balance:,} 원")


if __name__ == "__main__":
    main()
