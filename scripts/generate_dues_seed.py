"""Generate DuesInitialData.kt from 한우리 Excel workbook.

All members use HALF_YEARLY (반기별).
Excel columns: B=년도, C=반기구분, D=납부일자, E=금액
"""
from __future__ import annotations

import re
from collections import defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

EXCEL_PATH = Path(r"C:\Users\gyoum\Downloads\한우리 모임\2025_11대_한우리-R3_(2026-01-05).xlsx")
OUT_KT = Path(__file__).resolve().parents[1] / "app/src/main/kotlin/com/smartexpense/data/local/seed/DuesInitialData.kt"

MEMBER_NAMES = [
    "강대화", "김민석", "김성겸", "김성환", "박성남",
    "유락준", "유진호", "최낙선", "전상환", "최창국",
    "김영섭",
]

STANDARD_HALF_FEE = 150_000
STANDARD_YEAR_FEE = STANDARD_HALF_FEE * 2
SKIP_YEAR_LABELS = {"이월미납", "미납계", "입회비"}


def split_lines(value) -> list[str]:
    if value is None:
        return []
    text = str(value).replace("\r\n", "\n").replace("\r", "\n")
    return [line.strip() for line in text.split("\n") if line.strip()]


def norm_date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if re.match(r"\d{4}-\d{2}-\d{2}", text):
        return text[:10]
    if re.match(r"\d{4}\.\d{2}\.\d{2}", text):
        return text.replace(".", "-")
    return ""


def norm_amount(value) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).replace(",", "").strip()
    if not text or text in ("미납", "-", "—"):
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def is_unpaid_marker(amount_cell) -> bool:
    return str(amount_cell).strip() == "미납" if amount_cell is not None else False


def half_target_amount(amount_cell) -> int:
    if is_unpaid_marker(amount_cell):
        return STANDARD_HALF_FEE
    parts = [norm_amount(part) for part in split_lines(amount_cell)]
    parts = [part for part in parts if part > 0]
    if parts:
        return sum(parts)
    single = norm_amount(amount_cell)
    return single if single > 0 else STANDARD_HALF_FEE


def parse_half_row(pay_date_cell, amount_cell) -> dict:
    target = half_target_amount(amount_cell)
    if is_unpaid_marker(amount_cell):
        return {"payDate": "", "amount": target, "isPaid": False}

    dates = [norm_date(part) for part in split_lines(pay_date_cell)]
    dates = [part for part in dates if part]
    if not dates and pay_date_cell is not None:
        single_date = norm_date(pay_date_cell)
        if single_date:
            dates = [single_date]

    amounts = [norm_amount(part) for part in split_lines(amount_cell)]
    amounts = [part for part in amounts if part > 0]
    amount = sum(amounts) if amounts else norm_amount(amount_cell)
    pay_date = max(dates) if dates else ""

    if pay_date and amount > 0:
        return {"payDate": pay_date, "amount": amount, "isPaid": True}
    if amount > 0:
        return {"payDate": "", "amount": amount, "isPaid": False}
    return {"payDate": "", "amount": target, "isPaid": False}


def parse_member_half_rows(ws) -> list[dict]:
    rows = list(ws.iter_rows(values_only=True))
    header_idx = None
    for index, row in enumerate(rows):
        cells = [str(cell).strip() if cell is not None else "" for cell in row]
        if "년도" in cells and "반기구분" in cells:
            header_idx = index
            break
    if header_idx is None:
        return []

    header = [str(cell).strip() if cell is not None else "" for cell in rows[header_idx]]
    col = {name: index for index, name in enumerate(header) if name}
    half_rows: list[dict] = []
    current_year: int | None = None

    for row in rows[header_idx + 1 :]:
        if not row:
            continue

        year_raw = row[col["년도"]]
        half = row[col["반기구분"]]
        pay_date = row[col["납부일자"]]
        amount = row[col["금액"]]

        if year_raw is not None and str(year_raw).strip():
            year_text = str(year_raw).replace("년", "").strip()
            if year_text in SKIP_YEAR_LABELS:
                current_year = None
                continue
            if year_text.isdigit():
                current_year = int(year_text)

        if current_year is None:
            continue

        half_text = str(half).strip() if half else ""
        if "상" in half_text:
            term = "상반기"
        elif "하" in half_text:
            term = "하반기"
        else:
            continue

        parsed = parse_half_row(pay_date, amount)
        half_rows.append(
            {
                "year": current_year,
                "term": term,
                **parsed,
            }
        )

    return normalize_same_year_pairs(half_rows)


def normalize_same_year_pairs(half_rows: list[dict]) -> list[dict]:
    by_year: dict[int, dict[str, dict]] = defaultdict(dict)
    for row in half_rows:
        by_year[row["year"]][row["term"]] = dict(row)

    result: list[dict] = []
    for year in sorted(by_year.keys()):
        first = by_year[year].get("상반기")
        second = by_year[year].get("하반기")
        if not first or not second:
            if first:
                result.append(first)
            if second:
                result.append(second)
            continue

        # 연 회비를 상반기 한 줄에 모아 둔 경우 → 상·하 각각 반기 금액
        if (
            first["isPaid"]
            and first["amount"] >= STANDARD_YEAR_FEE
            and not second["isPaid"]
            and not second["payDate"]
        ):
            half_amount = first["amount"] // 2
            remainder = first["amount"] - half_amount
            first["amount"] = half_amount
            second["amount"] = remainder
            second["payDate"] = first["payDate"]
            second["isPaid"] = True

        # 날짜 병합: 상반기에만 날짜, 하반기에 금액만 있는 경우
        elif (
            not second["isPaid"]
            and second["amount"] > 0
            and not second["payDate"]
            and first["payDate"]
        ):
            second["payDate"] = first["payDate"]
            second["isPaid"] = True

        result.append(first)
        result.append(second)

    return result


def build_half_year_records(half_rows: list[dict]) -> list[dict]:
    return [
        {
            "year": row["year"],
            "term": row["term"],
            "amount": row["amount"],
            "payDate": row["payDate"],
            "isPaid": row["isPaid"],
        }
        for row in half_rows
    ]


def kotlin_record(record: dict) -> str:
    return (
        f'            DuesSeedRecord({record["year"]}, "{record["term"]}", '
        f'{record["amount"]}, "{record["payDate"]}", {str(record["isPaid"]).lower()})'
    )


def main() -> None:
    wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
    member_seeds: dict[str, list[dict]] = {}
    for name in MEMBER_NAMES:
        if name not in wb.sheetnames:
            raise SystemExit(f"Missing sheet: {name}")
        half_rows = parse_member_half_rows(wb[name])
        member_seeds[name] = build_half_year_records(half_rows)
    wb.close()

    # 원래 회원 순서 유지
    ordered_names = [
        "강대화", "김민석", "김성겸", "김성환", "김영섭", "박성남",
        "유락준", "유진호", "최낙선", "전상환", "최창국",
    ]

    lines = [
        "package com.smartexpense.data.local.seed",
        "",
        "import com.smartexpense.data.local.entity.club.DuesPaymentMethod",
        "",
        "/**",
        " * 엑셀 `2025_11대_한우리-R3_(2026-01-05).xlsx` 회원별 시트 (반기별 정리본).",
        " */",
        "object DuesInitialData {",
        "",
        "    data class DuesSeedRecord(",
        "        val year: Int,",
        "        val termLabel: String,",
        "        val amount: Int,",
        "        val payDate: String,",
        "        val isPaid: Boolean",
        "    )",
        "",
        "    data class MemberDuesSeed(",
        "        val paymentMethod: DuesPaymentMethod,",
        "        val records: List<DuesSeedRecord>",
        "    )",
        "",
        "    val seedsByMember: Map<String, MemberDuesSeed> = mapOf(",
    ]

    member_blocks = []
    for name in ordered_names:
        records = member_seeds[name]
        rec_lines = [kotlin_record(record) for record in records]
        member_blocks.append(
            f'        "{name}" to MemberDuesSeed(\n'
            f"            paymentMethod = DuesPaymentMethod.HALF_YEARLY,\n"
            f"            records = listOf(\n"
            + ",\n".join(rec_lines)
            + "\n            )\n        )"
        )
    lines.append(",\n".join(member_blocks))
    lines.extend(
        [
            "    )",
            "",
            "    fun totalTargetAmountForYear(records: List<DuesSeedRecord>, year: Int): Int =",
            "        records.filter { it.year == year }.sumOf { it.amount }",
            "}",
            "",
        ]
    )

    OUT_KT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Generated {OUT_KT}")
    for name in ordered_names:
        records = member_seeds[name]
        paid = sum(1 for record in records if record["isPaid"])
        unpaid_amt = sum(record["amount"] for record in records if not record["isPaid"])
        years = sorted({record["year"] for record in records})
        year_range = f"{years[0]}-{years[-1]}" if years else "-"
        print(
            f"  {name}: HALF_YEARLY, {len(records)} slots, "
            f"{paid} paid, unpaid {unpaid_amt:,}원, years {year_range}"
        )


if __name__ == "__main__":
    main()
