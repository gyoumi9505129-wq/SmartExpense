"""Generate TransactionInitialData.kt from 한우리 Excel ledger workbook."""
from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import openpyxl

EXCEL_PATH = Path(r"C:\Users\gyoum\Downloads\한우리거래내역.xlsx")
OUT_KT = (
    Path(__file__).resolve().parents[1]
    / "app/src/main/kotlin/com/smartexpense/data/local/seed/TransactionInitialData.kt"
)
# 초기 데이터 최종 잔액 기준 (수입 합 − 지출 합)
EXPECTED_FINAL_BALANCE = 14_794_878

REGULAR_DUES = "정기 회비 (월/연회비)"
DONATION = "찬조금/특별 회비 (기부금 등)"
OTHER_INCOME = "기타 수입 (이자, 이월금 등)"
MEAL = "식대/다과비 (모임 식사, 카페 등)"
VENUE_RENTAL = "장소 대관료 (모임 장소 대여비)"
EVENT = "행사/진행비 (이벤트 기획, 상품 구매 등)"
SUPPLIES = "비품/소모품비 (운영에 필요한 물품 구매)"
CONDOLENCE = "경조사비 (회원 화환, 축의금/조의금 등)"
TRANSPORT = "교통/통신비 (이동 경비, 문자 발송비 등)"
OTHER_EXPENSE = "기타 지출 (수수료 및 예비비)"

SKIP_LABELS = {"구   분", "구분", "기 결 산", "기결산"}


def sheet_sort_key(name: str) -> tuple[int, int]:
    match = re.match(r"(\d{4})년(?:-(\d))?", name)
    if not match:
        return 9999, 9
    return int(match.group(1)), int(match.group(2) or 0)


def norm_date(value) -> str:
    if value is None or value == "":
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    text = str(value).strip()
    if re.match(r"\d{4}-\d{2}-\d{2}", text):
        return text[:10]
    if re.match(r"\d{4}\.\d{2}\.\d{2}", text):
        return text.replace(".", "-")
    return ""


def norm_amount(value) -> int:
    if value is None or value == "":
        return 0
    if isinstance(value, (int, float)):
        return int(value)
    text = str(value).replace(",", "").strip()
    if not text or text in ("-", "—"):
        return 0
    try:
        return int(float(text))
    except ValueError:
        return 0


def norm_note(value) -> str:
    if value is None:
        return ""
    return str(value).strip()


def is_carryover_only(label: str, income: int, expense: int) -> bool:
    if income > 0 or expense > 0:
        return False
    return any(keyword in label for keyword in ("이월", "잔여금", "잔액", "기결산", "기 결산"))


def classify(label: str, income: int, expense: int, note: str) -> tuple[str, str]:
    text = f"{label} {note}".lower()

    if income > 0:
        if any(keyword in text for keyword in ("회비", "분납", "납부", "미납회비", "입회비")):
            return "INCOME", REGULAR_DUES
        if any(keyword in text for keyword in ("찬조", "기부", "후원")):
            return "INCOME", DONATION
        if any(
            keyword in text
            for keyword in (
                "이자",
                "이월",
                "잔여",
                "잔액",
                "예금이자",
                "되돌려",
                "반환",
                "환급",
                "전대",
            )
        ):
            return "INCOME", OTHER_INCOME
        if any(keyword in text for keyword in ("은행", "우체국", "농협", "신한")):
            return "INCOME", OTHER_INCOME
        if re.fullmatch(r"[가-힣]{2,5}", label.strip()):
            return "INCOME", REGULAR_DUES
        return "INCOME", OTHER_INCOME

    if any(
        keyword in text
        for keyword in ("경조", "화환", "꽃", "칠순", "고희", "조의", "축의", "부의", "출산", "결혼")
    ):
        return "EXPENSE", CONDOLENCE
    if any(keyword in text for keyword in ("펜션", "숙박", "숙소", "대관", "장소대관", "모텔")):
        return "EXPENSE", VENUE_RENTAL
    if any(
        keyword in text
        for keyword in (
            "모임",
            "식대",
            "식비",
            "다과",
            "부식",
            "주류",
            "음료",
            "가리비",
            "케잌",
            "케이크",
            "정산",
            "바베큐",
            "회식",
            "식사",
            "명절",
        )
    ):
        return "EXPENSE", MEAL
    if any(keyword in text for keyword in ("행사", "이벤트", "진행비", "상품", "게임", "경품", "레크레이션")):
        return "EXPENSE", EVENT
    if any(keyword in text for keyword in ("비품", "소모품", "운영비", "잡비", "용품", "구입비")):
        return "EXPENSE", SUPPLIES
    if any(keyword in text for keyword in ("주전부리",)):
        return "EXPENSE", MEAL
    if any(
        keyword in text
        for keyword in (
            "교통",
            "통신",
            "문자",
            "주유",
            "유류",
            "택시",
            "톨비",
            "송금수수료",
            "수수료",
            "통행",
            "주차",
        )
    ):
        return "EXPENSE", TRANSPORT
    return "EXPENSE", OTHER_EXPENSE


def build_note(label: str, remark: str) -> str:
    remark = remark.strip()
    label = label.strip()
    combined = f"{label} | {remark}" if remark and remark != label else label
    return " ".join(combined.split())


def parse_rows(ws) -> list[dict]:
    records: list[dict] = []
    for row in ws.iter_rows(values_only=True):
        if not row or len(row) < 5:
            continue
        label = str(row[0]).strip() if row[0] else ""
        if not label or label in SKIP_LABELS or label.startswith("◆"):
            continue

        pay_date = norm_date(row[1])
        income = norm_amount(row[2])
        expense = norm_amount(row[3])
        balance_after = norm_amount(row[4]) if len(row) > 4 else 0
        remark = norm_note(row[5]) if len(row) > 5 else ""

        if is_carryover_only(label, income, expense):
            continue
        if income == 0 and expense == 0:
            continue
        if not pay_date:
            continue

        tx_type, category = classify(label, income, expense, remark)
        records.append(
            {
                "date": pay_date,
                "type": tx_type,
                "category": category,
                "incomeAmount": income,
                "expenseAmount": expense,
                "note": build_note(label, remark),
                "balanceAfter": balance_after if balance_after != 0 else None,
            }
        )
    return records


def parse_ledger_workbook() -> list[dict]:
    wb = openpyxl.load_workbook(EXCEL_PATH, read_only=True, data_only=True)
    year_sheets = sorted(
        [name for name in wb.sheetnames if re.match(r"\d{4}년", name)],
        key=sheet_sort_key,
    )

    # 연도별 시트(구 형식) 또는 단일 통합 시트(신 형식) 모두 지원
    if year_sheets:
        records: list[dict] = []
        for sheet_name in year_sheets:
            records.extend(parse_rows(wb[sheet_name]))
    else:
        sheet_name = "Sheet1" if "Sheet1" in wb.sheetnames else wb.sheetnames[0]
        records = parse_rows(wb[sheet_name])

    wb.close()
    return records


def kotlin_escape(text: str) -> str:
    return (
        text.replace("\\", "\\\\")
        .replace('"', '\\"')
        .replace("\r\n", "\\n")
        .replace("\n", "\\n")
        .replace("\r", "\\n")
    )


def kotlin_record(record: dict) -> str:
    note = kotlin_escape(record["note"])
    balance = (
        "null"
        if record["balanceAfter"] is None
        else str(record["balanceAfter"])
    )
    return (
        f'            TransactionSeedRecord('
        f'"{record["date"]}", '
        f'ClubTransactionType.{record["type"]}, '
        f'"{record["category"]}", '
        f'{record["incomeAmount"]}, '
        f'{record["expenseAmount"]}, '
        f'"{note}", '
        f'{balance})'
    )


def main() -> None:
    records = parse_ledger_workbook()
    lines = [
        "package com.smartexpense.data.local.seed",
        "",
        "import com.smartexpense.data.local.entity.club.ClubTransactionType",
        "",
        "/**",
        " * 엑셀 `한우리거래내역.xlsx` 입출금 기록부 시드.",
        " */",
        "object TransactionInitialData {",
        "",
        "    data class TransactionSeedRecord(",
        "        val date: String,",
        "        val type: ClubTransactionType,",
        "        val category: String,",
        "        val incomeAmount: Int,",
        "        val expenseAmount: Int,",
        "        val note: String,",
        "        val balanceAfter: Int? = null,",
        "    )",
        "",
        "    val records: List<TransactionSeedRecord> = listOf(",
        ",\n".join(kotlin_record(record) for record in records),
        "    )",
        "}",
        "",
    ]
    OUT_KT.write_text("\n".join(lines), encoding="utf-8")
    print(f"Generated {OUT_KT} with {len(records)} records")
    income_total = sum(record["incomeAmount"] for record in records)
    expense_total = sum(record["expenseAmount"] for record in records)
    balance = income_total - expense_total
    print(f"  Income: {income_total:,}  Expense: {expense_total:,}  Balance: {balance:,}")
    if balance != EXPECTED_FINAL_BALANCE:
        raise SystemExit(
            f"최종 잔액 {balance:,} != 기대값 {EXPECTED_FINAL_BALANCE:,}. "
            "엑셀/파서를 확인하세요."
        )


if __name__ == "__main__":
    main()
